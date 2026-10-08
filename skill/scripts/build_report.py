#!/usr/bin/env python3
"""
build_report.py - build the aws-audit Word report from run_checks.py output.

Usage (from the skill/ directory, in the venv - see requirements.txt):
    python3 scripts/build_report.py ../audit-runs/<account_id>-<YYYYMMDD>/findings.json

    # A retest, labeled in the filename and on the cover:
    python3 scripts/build_report.py ../audit-runs/<id>-<date>/findings-retest-1.json \\
        --label retest-1 --variant-note "Retest 1 - post-remediation"

Inputs, all from the engagement directory (the findings file's directory
unless overridden):
  - findings.json        run_checks.py output (required, positional)
  - engagement-info.md   cover page, Prepared By block, Report Status and the
                         Testing Narrative table (required; --engagement-info
                         to override). Format: references/engagement-info-template.md.
  - snapshot.json        optional, for regions and the collecting identity in
                         the methodology appendix. Defaults to the findings
                         file's name with "findings" replaced by "snapshot",
                         if that file exists; --snapshot to override.

Output: AWS_Security_Audit_<account_id>_<framework>[_<label>]_<DRAFT|FINAL>.docx
in the engagement directory (or --out-dir). DRAFT vs FINAL comes only from
engagement-info.md's "Report Status" field and drives the title, cover
banner, page header and filename together - see
references/remediation-and-retesting.md, "Finalizing." --variant-note
overrides engagement-info.md's "Report Variant Note" for this build only.

Building a FINAL report prints a WARNING for anything the finalization
checklist says must be settled first (open "fail" findings, "Pending"
Testing Narrative rows, blank auditor fields). It still builds, so the
warnings can be reviewed against the output, but a FINAL with warnings
isn't ready to deliver.
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime, timezone

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("python-docx is required: activate a venv and run pip install -r requirements.txt (from the skill/ directory)", file=sys.stderr)
    sys.exit(1)


SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CHECK_CATALOG = os.path.join(SKILL_DIR, "references", "check-catalog.md")

FRAMEWORK_NAMES = {
    "cis": "CIS AWS Foundations Benchmark v7.0.0",
    "well_architected": "AWS Well-Architected Framework, Security Pillar",
    "soc2": "SOC 2 Trust Services Criteria (Common Criteria / Security)",
    "iso27001": "ISO/IEC 27001:2022 Annex A",
    "all": "All supported frameworks (CIS, Well-Architected, SOC 2, ISO 27001)",
}
FRAMEWORK_SHORT = {"cis": "CIS", "well_architected": "Well-Architected", "soc2": "SOC 2", "iso27001": "ISO 27001"}
AREA_ORDER = ["IAM", "MFA", "S3", "CloudTrail / Logging", "Security Groups / VPC", "Root Account"]
SEVERITY_ORDER = ["Critical", "High", "Medium", "Low"]
OUT_OF_SCOPE = (
    "RDS, EFS, EBS, KMS (general), AWS Config, AWS Organizations governance, "
    "Security Hub / GuardDuty enablement, and other services outside the six core areas."
)

METADATA_KEYS = [
    "Report Status", "Report Variant Note", "Auditor Name", "Auditor Title",
    "Auditor Organization", "Auditor Contact", "Account/Engagement Owner", "Primary Contact",
]
REQUIRED_MILESTONES = ["Test Start", "Test Completion", "Findings Draft", "Findings Review", "Findings Completion"]
NOT_PROVIDED = "Not provided"
PENDING = "Pending"

# Colors
DRAFT_FILL, FINAL_FILL = "C00000", "2E7D32"
HEADER_FILL = "1F3864"
SEVERITY_FILL = {"Critical": "C00000", "High": "E36C09", "Medium": "BF8F00", "Low": "548235"}
LIGHT_FILL = "F2F2F2"


class BuildError(Exception):
    pass


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

def _section(text, heading):
    """Return the body of a '## heading' section (up to the next '## ')."""
    m = re.search(rf"^##\s+{re.escape(heading)}\s*$(.*?)(?=^##\s|\Z)", text, re.M | re.S)
    return m.group(1) if m else None


def parse_engagement_info(path):
    """Parse the two machine-read sections of engagement-info.md. Anything
    outside them, and any line in them that isn't a known '- Key: Value'
    bullet or a table row, is ignored (the template's italic notes)."""
    if not os.path.exists(path):
        raise BuildError(
            f"{path} not found. Copy references/engagement-info-template.md there and fill it in "
            "(SKILL.md Step 1)."
        )
    with open(path) as f:
        text = f.read()

    meta_body = _section(text, "Report Metadata")
    if meta_body is None:
        raise BuildError(f"{path} has no '## Report Metadata' section.")
    meta = {}
    for line in meta_body.splitlines():
        m = re.match(r"^\s*-\s+([^:*`]+?):\s*(.*?)\s*$", line)
        if m and m.group(1) in METADATA_KEYS:
            meta[m.group(1)] = m.group(2)
    status = meta.get("Report Status", "").strip().upper()
    if status not in ("DRAFT", "FINAL"):
        raise BuildError(f"{path}: Report Status must be DRAFT or FINAL, got {meta.get('Report Status')!r}.")
    meta["Report Status"] = status

    narrative_body = _section(text, "Testing Narrative")
    if narrative_body is None:
        raise BuildError(f"{path} has no '## Testing Narrative' section.")
    milestones = []
    for line in narrative_body.splitlines():
        line = line.strip()
        if not line.startswith("|") or re.match(r"^\|[\s|:-]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells[0] == "Milestone" or not cells[0]:
            continue
        cells += [""] * (4 - len(cells))
        milestones.append({"milestone": cells[0], "date": cells[1], "by": cells[2], "notes": cells[3]})
    names = [m["milestone"] for m in milestones]
    missing = [m for m in REQUIRED_MILESTONES if m not in names]
    if missing:
        raise BuildError(f"{path}: Testing Narrative is missing row(s): {', '.join(missing)}. Keep all five, even if blank.")
    return meta, milestones


def parse_manual_checks(path=CHECK_CATALOG):
    """The 'Manual/extended checks' bullets from check-catalog.md, as plain text."""
    try:
        with open(path) as f:
            text = f.read()
    except OSError:
        return []
    body = _section(text, "Manual/extended checks (not automated in v1)") or ""
    items = []
    for line in body.splitlines():
        m = re.match(r"^\s*-\s+(.*)$", line)
        if m:
            items.append(m.group(1).replace("**", "").replace("`", ""))
    return items


def default_snapshot_path(findings_path):
    d, name = os.path.split(findings_path)
    if "findings" not in name:
        return None
    candidate = os.path.join(d, name.replace("findings", "snapshot", 1))
    return candidate if os.path.exists(candidate) else None


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------

def _shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def shade_cell(cell, fill):
    _shade(cell._tc.get_or_add_tcPr(), fill)


def set_cell_text(cell, text, bold=False, color=None, size=None):
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return run


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL, font_size=9):
    """A full-width grid table with a shaded header row. widths in inches."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, color="FFFFFF", size=font_size)
        shade_cell(cell, header_fill)
    _repeat_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def add_field(paragraph, instr):
    """Insert a field (e.g. PAGE) into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)


def banner(doc, text, fill):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _shade(p._p.get_or_add_pPr(), fill)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    return p


# ---------------------------------------------------------------------------
# Report content
# ---------------------------------------------------------------------------

def format_timestamp(value):
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return str(value)
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def refs_text(finding, framework):
    refs = finding.get("refs") or {}
    if framework == "all":
        parts = [f"{FRAMEWORK_SHORT.get(k, k)}: {v}" for k, v in refs.items() if v]
        return "; ".join(parts) or "—"
    return refs.get(framework) or "—"


def area_key(area):
    return AREA_ORDER.index(area) if area in AREA_ORDER else len(AREA_ORDER)


def severity_key(sev):
    return SEVERITY_ORDER.index(sev) if sev in SEVERITY_ORDER else len(SEVERITY_ORDER)


def group_by_check(findings):
    """{check_id: [findings]} preserving first-seen order."""
    grouped = {}
    for f in findings:
        grouped.setdefault(f["check_id"], []).append(f)
    return grouped


def build_cover(doc, ctx):
    meta, status = ctx["meta"], ctx["status"]
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"AWS Account Security Audit Report ({status})")
    run.bold = True
    run.font.size = Pt(24)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(FRAMEWORK_NAMES.get(ctx["framework"], ctx["framework"])).font.size = Pt(14)
    if ctx["variant_note"]:
        vp = doc.add_paragraph()
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vr = vp.add_run(ctx["variant_note"])
        vr.italic = True

    if status == "DRAFT":
        banner(doc, "DRAFT — NOT FOR DISTRIBUTION AS A FINAL REPORT", DRAFT_FILL)
    else:
        banner(doc, "FINAL", FINAL_FILL)
    doc.add_paragraph()

    s = ctx["summary"]
    rows = [
        ("Report Status", status),
        ("AWS Account ID", s.get("account_id") or "Unknown"),
        ("Framework", FRAMEWORK_NAMES.get(ctx["framework"], ctx["framework"])),
        ("Data Collected", ctx["collected_at"]),
        ("Report Generated", ctx["generated_at"]),
        ("Account/Engagement Owner", meta.get("Account/Engagement Owner") or NOT_PROVIDED),
        ("Primary Contact", meta.get("Primary Contact") or NOT_PROVIDED),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, bold=True, size=10)
        shade_cell(cells[0], LIGHT_FILL)
        run = set_cell_text(cells[1], value, size=10, bold=(label == "Report Status"))
        if label == "Report Status":
            run.font.color.rgb = RGBColor.from_string(DRAFT_FILL if status == "DRAFT" else FINAL_FILL)
    for row in table.rows:
        row.cells[0].width, row.cells[1].width = Inches(2.2), Inches(4.3)
    doc.add_paragraph()

    doc.add_heading("Prepared By", level=2)
    for key in ("Auditor Name", "Auditor Title", "Auditor Organization", "Auditor Contact"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run(f"{key.replace('Auditor ', '')}: ").bold = True
        p.add_run(meta.get(key) or NOT_PROVIDED)
    doc.add_paragraph()

    doc.add_heading("Testing Narrative", level=2)
    add_table(
        doc,
        ["Milestone", "Date", "By", "Notes"],
        [(m["milestone"], m["date"] or PENDING, m["by"], m["notes"]) for m in ctx["milestones"]],
        [1.6, 1.1, 1.4, 2.4],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_executive_summary(doc, ctx):
    s, findings = ctx["summary"], ctx["findings"]
    doc.add_heading("1. Executive Summary", level=1)
    failed = [f for f in findings if f["status"] == "fail"]
    by_sev = s.get("by_severity") or {}
    crit_high = by_sev.get("Critical", 0) + by_sev.get("High", 0)
    if not failed:
        posture = "No failing findings were identified in the evaluated controls."
    elif crit_high:
        posture = (
            f"{crit_high} critical or high-severity finding(s) require prompt remediation; "
            "they are listed in Key Findings below."
        )
    else:
        posture = "No critical or high-severity findings were identified; the remaining findings are medium or low severity."
    doc.add_paragraph(
        f"This report covers AWS account {s.get('account_id') or 'Unknown'}, evaluated against the "
        f"{FRAMEWORK_NAMES.get(ctx['framework'], ctx['framework'])} across six core areas: "
        f"{', '.join(AREA_ORDER)}. Data was collected {ctx['collected_at']}. {posture}"
    )
    rows = [("Failed findings", s.get("failed", 0))]
    rows += [(f"    {sev}", by_sev.get(sev, 0)) for sev in SEVERITY_ORDER]
    rows += [
        ("Passed findings", s.get("passed", 0)),
        ("Accepted findings", s.get("accepted", 0)),
        ("Checks evaluated", len(s.get("checks_evaluated") or [])),
        ("Checks skipped (missing data)", len(s.get("checks_skipped") or [])),
    ]
    add_table(doc, ["Measure", "Count"], rows, [3.5, 1.5], font_size=10)


def build_key_findings(doc, ctx):
    doc.add_heading("2. Key Findings", level=1)
    key = [f for f in ctx["findings"] if f["status"] == "fail" and f["severity"] in ("Critical", "High")]
    if not key:
        doc.add_paragraph("No critical or high-severity findings.")
        return
    doc.add_paragraph("Critical and high-severity failing findings, most severe first, grouped by area.")
    key.sort(key=lambda f: (severity_key(f["severity"]), area_key(f["area"])))
    for sev in ("Critical", "High"):
        sev_findings = [f for f in key if f["severity"] == sev]
        if not sev_findings:
            continue
        doc.add_heading(sev, level=2)
        for area in sorted({f["area"] for f in sev_findings}, key=area_key):
            doc.add_heading(area, level=3)
            for fs in group_by_check(x for x in sev_findings if x["area"] == area).values():
                resources = ", ".join(str(f["resource"]) for f in fs)
                add_bullet(doc, f" — {resources} ({refs_text(fs[0], ctx['framework'])})", bold_prefix=fs[0]["title"])


def build_accepted(doc, ctx, number):
    accepted = [f for f in ctx["findings"] if f["status"] == "accepted"]
    doc.add_heading(f"{number}. Accepted Findings", level=1)
    doc.add_paragraph(
        "These findings are still true in the collected data, but have been reviewed and explicitly "
        "risk-accepted through the engagement's exceptions file. They are neither open issues nor "
        "passing controls. Each entry shows who accepted it, why, and the rule that matched."
    )
    accepted.sort(key=lambda f: (severity_key(f["severity"]), area_key(f["area"])))
    rows = []
    for f in accepted:
        # run_checks.py appends " Accepted: <note> (...)" to the original evidence.
        evidence, sep, acceptance = f["evidence"].partition(" Accepted: ")
        rows.append((f["title"], f["severity"], f["resource"], evidence if sep else "", acceptance if sep else f["evidence"]))
    add_table(
        doc,
        ["Finding", "Severity", "Resource", "Evidence", "Accepted because (from exceptions file)"],
        rows,
        [1.3, 0.7, 1.1, 1.5, 1.9],
        font_size=8,
    )


def build_passing(doc, ctx, number):
    doc.add_heading(f"{number}. Passing Controls", level=1)
    grouped = group_by_check(ctx["findings"])
    passing = [(fs[0], len(fs)) for fs in grouped.values() if all(f["status"] == "pass" for f in fs)]
    if not passing:
        doc.add_paragraph("No control passed for every evaluated resource.")
        return
    doc.add_paragraph("Controls that passed for every evaluated resource:")
    passing.sort(key=lambda x: (area_key(x[0]["area"]), x[0]["title"]))
    for f, n in passing:
        suffix = f" ({n} resources)" if n > 1 else ""
        add_bullet(doc, f" — {refs_text(f, ctx['framework'])}{suffix}", bold_prefix=f"{f['area']}: {f['title']}")


def build_detailed(doc, ctx, number):
    doc.add_heading(f"{number}. Detailed Findings by Area", level=1)
    doc.add_paragraph(
        "Every failing finding, by area, with its control reference, evidence and remediation. "
        "Accepted findings are listed in their own section."
    )
    failed = [f for f in ctx["findings"] if f["status"] == "fail"]
    for i, area in enumerate(AREA_ORDER + sorted({f["area"] for f in failed} - set(AREA_ORDER)), 1):
        area_failed = [f for f in failed if f["area"] == area]
        doc.add_heading(f"{number}.{i} {area}", level=2)
        if not area_failed:
            doc.add_paragraph("No failing findings in this area.")
            continue
        grouped = group_by_check(area_failed)
        for check_id in sorted(grouped, key=lambda c: (severity_key(grouped[c][0]["severity"]), c)):
            fs = grouped[check_id]
            f0 = fs[0]
            doc.add_heading(f0["title"], level=3)
            p = doc.add_paragraph()
            p.add_run("Severity: ").bold = True
            sev = p.add_run(f0["severity"])
            sev.bold = True
            sev.font.color.rgb = RGBColor.from_string(SEVERITY_FILL.get(f0["severity"], "000000"))
            p.add_run("    Control: ").bold = True
            p.add_run(refs_text(f0, ctx["framework"]))
            p.add_run("    Check ID: ").bold = True
            p.add_run(check_id)
            add_table(doc, ["Resource", "Evidence"], [(f["resource"], f["evidence"]) for f in fs], [1.8, 4.7])
            rp = doc.add_paragraph()
            rp.add_run("Remediation: ").bold = True
            rp.add_run(f0["remediation"])


def build_appendix(doc, ctx):
    s = ctx["summary"]
    doc.add_heading("Appendix: Scope & Methodology", level=1)

    doc.add_heading("Scope", level=2)
    doc.add_paragraph(f"In scope: {', '.join(AREA_ORDER)}.")
    doc.add_paragraph(f"Out of scope for this version: {OUT_OF_SCOPE}")

    doc.add_heading("Data source", level=2)
    snap = ctx["snapshot"] or {}
    add_bullet(doc, ctx["data_source"], bold_prefix="Source: ")
    add_bullet(doc, ctx["collected_at"], bold_prefix="Collected: ")
    if snap.get("caller_arn"):
        add_bullet(doc, snap["caller_arn"], bold_prefix="Collected as: ")
    if snap.get("regions"):
        add_bullet(doc, ", ".join(snap["regions"]), bold_prefix="Regions: ")
    add_bullet(doc, os.path.basename(ctx["findings_path"]), bold_prefix="Findings file: ")

    doc.add_heading("Controls evaluated", level=2)
    titles = {f["check_id"]: f["title"] for f in ctx["findings"]}
    evaluated = s.get("checks_evaluated") or []
    skipped = s.get("checks_skipped") or []
    doc.add_paragraph(f"{len(evaluated)} automated check(s) evaluated, {len(skipped)} skipped because the snapshot lacked the data they need.")
    add_table(
        doc,
        ["Check ID", "Title", "Result"],
        [(c, titles.get(c, "(no resources evaluated)"), "Evaluated") for c in evaluated]
        + [(c, "—", "Skipped: missing data") for c in skipped],
        [2.4, 3.1, 1.0],
        font_size=8,
    )

    manual = ctx["manual_checks"]
    if manual:
        doc.add_heading("Manual/extended checks (not automated)", level=2)
        doc.add_paragraph(
            "These framework controls are not evaluated by the automated checks in this version and are "
            "out of scope for this report:"
        )
        for item in manual:
            add_bullet(doc, item)

    doc.add_heading("Exceptions", level=2)
    if s.get("exceptions_applied"):
        doc.add_paragraph(
            "An exceptions file was applied when the checks were run. It is kept in this engagement's "
            "evidence package alongside the snapshot and findings, together with the documentation "
            "request response and supporting evidence it is based on."
        )
        warnings = s.get("exceptions_warnings") or []
        if warnings:
            doc.add_paragraph("run_checks.py reported these warnings about the exceptions file:")
            for w in warnings:
                add_bullet(doc, w)
    else:
        doc.add_paragraph("No exceptions file was applied.")


def setup_document(ctx):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run(f"{ctx['status']} — AWS Security Audit — Account {ctx['summary'].get('account_id') or 'Unknown'}")
    hr.font.size = Pt(8)
    hr.bold = ctx["status"] == "DRAFT"
    if ctx["status"] == "DRAFT":
        hr.font.color.rgb = RGBColor.from_string(DRAFT_FILL)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ").font.size = Pt(8)
    add_field(footer, "PAGE")

    props = doc.core_properties
    props.title = f"AWS Account Security Audit Report ({ctx['status']})"
    props.subject = FRAMEWORK_NAMES.get(ctx["framework"], ctx["framework"])
    props.author = ctx["meta"].get("Auditor Name") or ""
    return doc


def finalization_warnings(ctx):
    if ctx["status"] != "FINAL":
        return []
    warnings = []
    open_count = sum(1 for f in ctx["findings"] if f["status"] == "fail")
    if open_count:
        warnings.append(
            f"{open_count} finding(s) are still \"fail\". Every finding must be fixed-and-retested or accepted "
            "before FINAL (remediation-and-retesting.md, Finalizing step 1)."
        )
    pending = [m["milestone"] for m in ctx["milestones"] if not m["date"]]
    if pending:
        warnings.append(f"Testing Narrative row(s) still Pending: {', '.join(pending)}.")
    blank = [k for k in ("Auditor Name", "Auditor Title", "Auditor Organization", "Auditor Contact") if not ctx["meta"].get(k)]
    if blank:
        warnings.append(f"Blank auditor field(s) will print as \"{NOT_PROVIDED}\": {', '.join(blank)}.")
    return warnings


def output_filename(account_id, framework, label, status):
    parts = ["AWS_Security_Audit", account_id or "unknown", framework]
    if label:
        parts.append(re.sub(r"[^A-Za-z0-9._-]+", "-", label))
    parts.append(status)
    return "_".join(parts) + ".docx"


def build(args):
    with open(args.findings) as f:
        result = json.load(f)
    summary, findings = result.get("summary") or {}, result.get("findings") or []
    framework = summary.get("framework") or "cis"
    engagement_dir = os.path.dirname(os.path.abspath(args.findings))

    meta, milestones = parse_engagement_info(args.engagement_info or os.path.join(engagement_dir, "engagement-info.md"))

    snapshot_path = args.snapshot or default_snapshot_path(args.findings)
    snapshot = None
    if snapshot_path:
        with open(snapshot_path) as f:
            snapshot = json.load(f)

    ctx = {
        "summary": summary,
        "findings": findings,
        "framework": framework,
        "meta": meta,
        "milestones": milestones,
        "status": meta["Report Status"],
        "variant_note": args.variant_note if args.variant_note is not None else meta.get("Report Variant Note", ""),
        "collected_at": format_timestamp(summary["collected_at"]) if summary.get("collected_at") else "unknown",
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "snapshot": snapshot,
        "data_source": args.data_source,
        "findings_path": args.findings,
        "manual_checks": parse_manual_checks(),
    }

    doc = setup_document(ctx)
    build_cover(doc, ctx)
    build_executive_summary(doc, ctx)
    build_key_findings(doc, ctx)
    number = 3
    if any(f["status"] == "accepted" for f in findings):
        build_accepted(doc, ctx, number)
        number += 1
    build_passing(doc, ctx, number)
    build_detailed(doc, ctx, number + 1)
    build_appendix(doc, ctx)

    out_dir = args.out_dir or engagement_dir
    out_path = os.path.join(out_dir, output_filename(summary.get("account_id"), framework, args.label, ctx["status"]))
    doc.save(out_path)
    return out_path, finalization_warnings(ctx)


def main():
    parser = argparse.ArgumentParser(description="Build the aws-audit Word report from a run_checks.py findings file.")
    parser.add_argument("findings", help="findings JSON produced by run_checks.py")
    parser.add_argument("--engagement-info", help="engagement-info.md (default: next to the findings file)")
    parser.add_argument("--snapshot", help="snapshot JSON, for regions/collector identity in the appendix "
                                           "(default: the findings file's name with 'findings' -> 'snapshot', if present)")
    parser.add_argument("--label", help="extra filename label, e.g. retest-1")
    parser.add_argument("--variant-note", help="overrides engagement-info.md's Report Variant Note for this build")
    parser.add_argument("--data-source", default="Live AWS API collection via collect_aws_data.py (read-only audit role)",
                        help="how the snapshot was produced, for the methodology appendix")
    parser.add_argument("--out-dir", help="output directory (default: the findings file's directory)")
    args = parser.parse_args()

    try:
        out_path, warnings = build(args)
    except BuildError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
    print(f"Wrote {out_path}", file=sys.stderr)
    for w in warnings:
        print(f"WARNING: {w}", file=sys.stderr)


if __name__ == "__main__":
    main()
